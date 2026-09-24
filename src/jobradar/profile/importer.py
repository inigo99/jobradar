"""Reading an existing CV into a structured profile.

An honest warning, repeated in the README: reproducing an arbitrary PDF layout
is not what happens here, and any tool that claims otherwise is overselling.
What happens is that the *content* is extracted into
:class:`~jobradar.models.Profile`, and the CV is then re-rendered with a
template chosen to resemble the original. That is the trade worth making: a
structured profile can be tailored, scored and validated per job, whereas a
pixel-perfect copy can only be reprinted.

Two import paths:

*With a language model* — the CV text is parsed into the full structure,
including a first pass at the evidence and ceiling maps. Good enough that most
people only correct a couple of fields.

*Without one* — a section-and-bullet heuristic produces a skeleton with the
contact details, the section blocks and the bullets it could find. It is a
starting point the user finishes in the dashboard, not a finished profile, and
the onboarding wizard says so.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from ..errors import MissingDependencyError, ProfileError, describe_os_error
from ..llm import LLMClient
from ..llm.prompts import parse_cv
from ..models import (
    Bullet,
    Certification,
    Contact,
    Education,
    Experience,
    LanguageSkill,
    Profile,
    SkillGroup,
)
from ..taxonomy import label_for
from ..textutils import detect_language, slugify, strip_html
from . import vocabulary

log = logging.getLogger(__name__)

EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
PHONE = re.compile(r"(?:\+\d{1,3}[ .-]?)?(?:\(?\d{2,4}\)?[ .-]?){2,4}\d{2,4}")
URL = re.compile(r"(?:https?://)?(?:www\.)?((?:linkedin\.com|github\.com)/[\w\-/%.]+)", re.I)
BULLET_LINE = re.compile(r"^\s*(?:[•▪◦·*\-–—]|\d+[.)])\s+(.{15,})$")
# Matches "2021 - 2023", "2022-03 - present", "2019-09 – 2020-07" and the
# Spanish, French and German words for "present".
DATE_RANGE = re.compile(
    r"((?:19|20)\d{2})(?:-(\d{1,2}))?\s*(?:[-–—/]|to|hasta|bis|à)\s*"
    r"((?:19|20)\d{2}(?:-\d{1,2})?|present|actualidad|actualmente|current|now|heute|présent)",
    re.I,
)

#: Section headings, in the languages CVs are most often written in.
SECTION_HEADINGS: dict[str, tuple[str, ...]] = {
    "summary": ("professional summary", "summary", "profile", "about me", "perfil",
                "resumen profesional", "resumen", "profil", "sobre mi"),
    "experience": ("professional experience", "work experience", "experience", "employment",
                   "experiencia profesional", "experiencia laboral", "experiencia",
                   "expérience professionnelle", "berufserfahrung"),
    "education": ("education", "academic background", "formación académica", "formacion academica",
                  "formación", "estudios", "formation", "ausbildung"),
    "certifications": ("certifications", "courses", "additional training", "certificaciones",
                       "cursos", "formación complementaria", "zertifikate"),
    "skills": ("technical skills", "skills", "competencies", "competencias técnicas",
               "competencias", "habilidades", "compétences", "kenntnisse"),
    "languages": ("languages", "idiomas", "langues", "sprachen"),
}


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------


#: Formats people commonly try that cannot be read, with what to do instead.
UNSUPPORTED_SUFFIXES = {
    ".doc": "Old Word .doc files cannot be read; save it as .docx or PDF.",
    ".odt": "OpenDocument files cannot be read; export it as PDF or .docx.",
    ".pages": "Pages files cannot be read; export it as PDF or .docx.",
    ".png": "Images cannot be read; upload the PDF or Word version of your CV.",
    ".jpg": "Images cannot be read; upload the PDF or Word version of your CV.",
    ".jpeg": "Images cannot be read; upload the PDF or Word version of your CV.",
}


def extract_text(path: str | Path) -> str:
    """Plain text from a PDF, DOCX, HTML, Markdown or text CV.

    The optional parsers are imported lazily so that a user who only ever
    uploads a ``.txt`` CV never has to install them. Every failure is raised
    as a :class:`ProfileError` (or :class:`MissingDependencyError`) that says
    what to do about it.
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if not path.exists():
        raise ProfileError(f"The CV file {path} does not exist.", hint="Check the path.")
    if path.is_dir():
        raise ProfileError(f"{path} is a folder, not a CV file.", hint="Pass the file itself.")
    if suffix in UNSUPPORTED_SUFFIXES:
        raise ProfileError(f"Cannot read {path.name}.", hint=UNSUPPORTED_SUFFIXES[suffix])

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            from pypdf.errors import PyPdfError
        except ImportError as exc:
            raise MissingDependencyError(
                "Reading PDF CVs needs the 'parse' extra.",
                hint="pip install 'jobradar[parse]'",
            ) from exc
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise ProfileError(
                    f"{path.name} is password-protected.",
                    hint="Save an unprotected copy of the PDF and upload that.",
                )
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ProfileError:
            raise
        except (PyPdfError, ValueError, KeyError, TypeError) as exc:
            raise ProfileError(
                f"{path.name} could not be read as a PDF: {exc}",
                hint="Re-export it from your editor, or upload the .docx version.",
            ) from exc
        except OSError as exc:
            raise ProfileError(f"Cannot read {path}: {describe_os_error(exc)}.") from exc

    if suffix in (".docx", ".dotx"):
        try:
            import docx
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as exc:
            raise MissingDependencyError(
                "Reading Word CVs needs the 'parse' extra.",
                hint="pip install 'jobradar[parse]'",
            ) from exc
        try:
            document = docx.Document(str(path))
        except (PackageNotFoundError, ValueError, KeyError) as exc:
            raise ProfileError(
                f"{path.name} could not be read as a Word document.",
                hint="Open it in Word and save it again as .docx, or upload a PDF.",
            ) from exc
        except OSError as exc:
            raise ProfileError(f"Cannot read {path}: {describe_os_error(exc)}.") from exc
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise ProfileError(f"Cannot read {path}: {describe_os_error(exc)}.") from exc
    if b"\x00" in raw[:4096]:
        raise ProfileError(
            f"{path.name} is not a text file.",
            hint="Supported formats: PDF, DOCX, HTML, Markdown or plain text.",
        )
    text = raw.decode("utf-8", errors="replace")
    return strip_html(text) if suffix in (".html", ".htm") else text


# ---------------------------------------------------------------------------
# Heuristic parsing (the no-model path)
# ---------------------------------------------------------------------------


def _classify_heading(line: str) -> str | None:
    """Is this line a section heading, and if so which section?

    Matching is near-exact on purpose. Prefix matching looks more forgiving and
    is actively harmful: "Languages: Python, Go, SQL" inside a skills block
    starts with "languages", and treating it as a heading silently swallows the
    whole skills section.
    """
    candidate = line.strip().rstrip(":").strip().lower()
    if not candidate or len(candidate) > 45 or candidate.endswith("."):
        return None
    # A heading is a label, not a sentence carrying values.
    if ":" in candidate or any(ch.isdigit() for ch in candidate):
        return None
    if len(candidate.split()) > 4:
        return None
    for section, headings in SECTION_HEADINGS.items():
        if candidate in headings:
            return section
    return None


def split_sections(text: str) -> dict[str, list[str]]:
    """Split CV text into the sections it declares.

    Anything before the first recognised heading becomes ``header`` — that is
    where the name and contact details live in essentially every CV layout.
    """
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        heading = _classify_heading(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _parse_contact(header: list[str], language: str = "en") -> Contact:
    blob = "\n".join(header)
    contact = Contact()
    email = EMAIL.search(blob)
    if email:
        contact.email = email.group(0)
    phone = PHONE.search(blob.replace(contact.email or "", " "))
    if phone and len(re.sub(r"\D", "", phone.group(0))) >= 8:
        contact.phone = phone.group(0).strip()
    for match in URL.finditer(blob):
        link = match.group(1)
        if "linkedin" in link.lower():
            contact.linkedin = link
        elif "github" in link.lower():
            contact.github = link
    # The name is normally the first line, and normally the only line with no
    # digits and no "@".
    for line in header[:4]:
        stripped = line.strip()
        if stripped and not any(ch.isdigit() for ch in stripped) and "@" not in stripped:
            if 2 <= len(stripped.split()) <= 6:
                contact.full_name = stripped.title() if stripped.isupper() else stripped
                break

    # The location is the "City, Country" fragment of the contact line — the
    # part that is not the name, the email, the phone or a URL.
    for line in header[:4]:
        for chunk in re.split(r"\s*[·|]\s*", line):
            chunk = chunk.strip()
            if not chunk or chunk == contact.full_name or "@" in chunk or "/" in chunk:
                continue
            if any(ch.isdigit() for ch in chunk) or "," not in chunk:
                continue
            if 3 < len(chunk) < 60:
                contact.city = {language: chunk}
                return contact
    return contact


def _parse_experience(lines: list[str], language: str) -> list[Experience]:
    """Group the experience section into positions and their bullets.

    Real CVs spread a position's header over two or three lines — title and
    employer on one, city and dates on the next. So a non-bullet line only
    starts a *new* position once the current one has collected at least one
    achievement; until then it is absorbed into the header being built.
    """
    experiences: list[Experience] = []
    current: Experience | None = None

    for line in lines:
        bullet_match = BULLET_LINE.match(line)
        if bullet_match:
            if current is not None:
                text = bullet_match.group(1).strip()
                current.bullets.append(
                    Bullet(id=f"{current.id}-{len(current.bullets) + 1}", text={language: text})
                )
            continue

        stripped = line.strip()
        if not stripped:
            continue

        if current is not None and not current.bullets:
            _absorb_header_line(current, stripped, language)
            continue

        current = _new_experience(stripped, language, len(experiences))
        experiences.append(current)

    return experiences


def _new_experience(line: str, language: str, index: int) -> Experience:
    """Start a position from its first header line."""
    title = re.split(r"\s+[—–|]\s+|,\s{2,}", line)[0].strip()
    identifier = slugify(title, 24) or f"role{index + 1}"
    experience = Experience(
        id=f"{identifier}-{index + 1}",
        title={language: title},
        organization=_organisation_from(line),
    )
    _apply_dates(experience, line)
    return experience


def _absorb_header_line(experience: Experience, line: str, language: str) -> None:
    """Fold a second or third header line into the position being built."""
    _apply_dates(experience, line)
    if not experience.organization:
        experience.organization = _organisation_from(line)
    if not experience.location:
        without_dates = DATE_RANGE.sub("", line).strip(" ·|—–-,")
        if without_dates:
            experience.location = {language: without_dates}


def _apply_dates(experience: Experience, line: str) -> None:
    """Set start and end from a date range found on a header line."""
    dates = DATE_RANGE.search(line)
    if not dates:
        return
    year, month, end = dates.group(1), dates.group(2), dates.group(3)
    experience.start = experience.start or f"{year}-{int(month or 1):02d}"
    if re.match(r"(?i)present|actualidad|actualmente|current|now|heute|présent", end):
        experience.end = None
    elif "-" in end:
        end_year, end_month = end.split("-")
        experience.end = f"{end_year}-{int(end_month):02d}"
    else:
        experience.end = f"{end}-12"


def _apply_dates_to_education(entry: Education, dates: re.Match[str]) -> None:
    year, month, end = dates.group(1), dates.group(2), dates.group(3)
    entry.start = f"{year}-{int(month or 1):02d}"
    if "-" in end:
        end_year, end_month = end.split("-")
        entry.end = f"{end_year}-{int(end_month):02d}"
    elif end.isdigit():
        entry.end = f"{end}-12"


def _organisation_from(line: str) -> str:
    parts = re.split(r"\s+[—–|]\s+|\sat\s|\sen\s", line.strip())
    return parts[1].strip() if len(parts) > 1 else ""


def _parse_education(lines: list[str], language: str) -> list[Education]:
    education: list[Education] = []
    for index, line in enumerate(lines):
        bullet = BULLET_LINE.match(line)
        if bullet and education:
            education[-1].note = {language: bullet.group(1)}
            continue
        dates = DATE_RANGE.search(line)
        degree = re.split(r"\s+[—–|]\s+", line.strip())[0].strip()
        if len(degree) < 5:
            continue
        entry = Education(
            id=f"edu-{index + 1}",
            degree={language: degree},
            institution={language: _organisation_from(line)},
        )
        if dates:
            _apply_dates_to_education(entry, dates)
        education.append(entry)
    return education


def _parse_skills(lines: list[str], language: str) -> list[SkillGroup]:
    groups: list[SkillGroup] = []
    for index, line in enumerate(lines):
        bullet = BULLET_LINE.match(line)
        text = bullet.group(1) if bullet else line.strip()
        label, separator, items = text.partition(":")
        if not separator:
            label, items = f"Skills {index + 1}", text
        parsed = [item.strip() for item in re.split(r"[,;·/|]", items) if item.strip()]
        if parsed:
            groups.append(
                SkillGroup(key=slugify(label, 20) or f"group{index}",
                           label={language: label.strip()}, items=parsed)
            )
    return groups


def heuristic_profile(text: str) -> Profile:
    """A best-effort profile with no language model involved."""
    language = detect_language(text)
    sections = split_sections(text)
    profile = Profile(default_language=language)
    profile.contact = _parse_contact(sections.get("header", []), language)
    if sections.get("summary"):
        profile.summary = {language: " ".join(sections["summary"])}
    profile.experience = _parse_experience(sections.get("experience", []), language)
    profile.education = _parse_education(sections.get("education", []), language)
    profile.skills = _parse_skills(sections.get("skills", []), language)
    for line in sections.get("certifications", []):
        bullet = BULLET_LINE.match(line)
        clean = bullet.group(1) if bullet else line.strip()
        year = re.search(r"(19|20)\d{2}", clean)
        profile.certifications.append(
            Certification(name={language: clean}, year=year.group(0) if year else "")
        )
    for line in sections.get("languages", []):
        for chunk in re.split(r"[·,;|]", line):
            chunk = chunk.strip()
            if not chunk:
                continue
            level = re.search(r"\(([^)]+)\)|\b([ABC][12])\b|native|nativo|fluent", chunk, re.I)
            profile.languages.append(
                LanguageSkill(
                    name={language: re.sub(r"\(.*?\)", "", chunk).strip()},
                    level=level.group(0).strip("()") if level else "",
                )
            )
    return vocabulary.refresh(profile)


# ---------------------------------------------------------------------------
# Model-assisted parsing
# ---------------------------------------------------------------------------


def _profile_from_payload(payload: dict, language: str) -> Profile:
    """Build a Profile from the JSON the model returned, defensively."""
    profile = Profile(default_language=language)
    contact = payload.get("contact") or {}
    profile.contact = Contact(
        full_name=str(contact.get("full_name", "")),
        city={language: str(contact.get("city", ""))},
        country=str(contact.get("country", "") or "").upper()[:2],
        phone=str(contact.get("phone", "")),
        email=str(contact.get("email", "")),
        linkedin=str(contact.get("linkedin", "")),
        github=str(contact.get("github", "")),
        website=str(contact.get("website", "")),
    )
    if payload.get("summary"):
        profile.summary = {language: str(payload["summary"])}

    for index, entry in enumerate(payload.get("experience") or []):
        experience = Experience(
            id=str(entry.get("id") or f"role-{index + 1}"),
            title={language: str(entry.get("title", ""))},
            organization=str(entry.get("organization", "")),
            location={language: str(entry.get("location", ""))},
            start=str(entry.get("start") or ""),
            end=str(entry["end"]) if entry.get("end") else None,
        )
        for bullet_index, bullet in enumerate(entry.get("bullets") or []):
            experience.bullets.append(
                Bullet(
                    id=str(bullet.get("id") or f"{experience.id}-{bullet_index + 1}"),
                    text={language: str(bullet.get("text", ""))},
                    skills=[str(s) for s in (bullet.get("skills") or [])],
                )
            )
        profile.experience.append(experience)

    for index, entry in enumerate(payload.get("education") or []):
        profile.education.append(
            Education(
                id=str(entry.get("id") or f"edu-{index + 1}"),
                degree={language: str(entry.get("degree", ""))},
                institution={language: str(entry.get("institution", ""))},
                start=str(entry.get("start") or ""),
                end=str(entry.get("end") or ""),
                note={language: str(entry.get("note", ""))},
            )
        )

    for entry in payload.get("certifications") or []:
        profile.certifications.append(
            Certification(
                name={language: str(entry.get("name", ""))},
                issuer=str(entry.get("issuer", "")),
                year=str(entry.get("year", "")),
            )
        )

    for index, entry in enumerate(payload.get("skills") or []):
        profile.skills.append(
            SkillGroup(
                key=str(entry.get("key") or f"group-{index + 1}"),
                label={language: str(entry.get("label", ""))},
                items=[str(item) for item in (entry.get("items") or [])],
            )
        )

    for entry in payload.get("languages") or []:
        profile.languages.append(
            LanguageSkill(name={language: str(entry.get("name", ""))}, level=str(entry.get("level", "")))
        )

    evidence = payload.get("evidence")
    if isinstance(evidence, dict):
        profile.evidence = {str(k): float(v) for k, v in evidence.items() if _is_number(v)}
    ceiling = payload.get("ceiling")
    if isinstance(ceiling, dict):
        profile.ceiling = {str(k): float(v) for k, v in ceiling.items() if _is_number(v)}
    return profile


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def import_profile(source: str | Path, llm: LLMClient | None = None) -> tuple[Profile, list[str]]:
    """Import a CV from a file (or raw text) into a profile.

    Returns the profile plus a list of notes for the user — what could not be
    determined and should be checked by hand. Onboarding shows these, because a
    silently wrong profile poisons every document generated afterwards.
    """
    text = str(source)
    if isinstance(source, Path):
        text = extract_text(source)
    elif "\n" not in text and len(text) < 400:
        try:
            candidate = Path(text)
            exists = candidate.exists()
        except (OSError, ValueError):  # not a usable path — treat it as raw CV text
            exists = False
        if exists:
            text = extract_text(candidate)
        elif candidate.suffix.lower() in (".pdf", ".docx", ".dotx", *UNSUPPORTED_SUFFIXES):
            # Clearly meant as a file name, not as the text of a CV.
            raise ProfileError(f"The CV file {text} does not exist.", hint="Check the path.")

    if not text.strip():
        raise ProfileError(
            "No text could be read from the CV.",
            hint="If it is a scanned PDF, export a text PDF or paste the text instead.",
        )

    notes: list[str] = []
    if len(text.strip()) < 200:
        notes.append("Very little text could be read from that file — check it is not a scan.")

    language = detect_language(text)
    profile: Profile | None = None

    if llm:
        system, user = parse_cv(text, language)
        payload = llm.complete_json(system, user, max_tokens=6000)
        if isinstance(payload, dict):
            profile = _profile_from_payload(payload, language)
        else:
            notes.append("The language model could not parse the CV; fell back to the text parser.")

    if profile is None:
        profile = heuristic_profile(text)
        notes.append(
            "Parsed without a language model: review the experience, dates and skills, "
            "which the text parser gets approximately right at best."
        )

    # Evidence is always derived from the text; a model's own numbers only
    # refine skills that derivation recognises (see merge_model_judgement).
    profile = vocabulary.merge_model_judgement(
        profile, dict(profile.evidence or {}), dict(profile.ceiling or {})
    )
    # Labels always come from the taxonomy so the dashboard shows "PyTorch /
    # TensorFlow" rather than the raw key.
    profile.skill_labels = {key: label_for(key) for key in (profile.evidence or {})}

    contact = profile.contact
    if not contact or not contact.full_name:
        notes.append("Could not find your name in the CV — add it in settings.")
    if not profile.experience:
        notes.append("No work experience was recognised — add at least one position.")
    for experience in (profile.experience or []):
        if not (experience.bullets or []):
            notes.append(f"'{experience.organization or experience.id}' has no achievements yet.")
    return profile, notes


def profile_from_form(data: dict, language: str = "en") -> Profile:
    """Build a minimal profile from the onboarding form, with no CV file.

    The wizard offers this for people who would rather type than upload; the
    profile is intentionally sparse and the dashboard nudges them to add
    achievements, because a CV with no measured results cannot be tailored into
    anything worth sending.
    """
    profile = Profile(default_language=language)
    profile.contact = Contact(
        full_name=str(data.get("full_name", "")),
        city={language: str(data.get("city", ""))},
        country=str(data.get("country", "")).upper()[:2],
        email=str(data.get("email", "")),
        phone=str(data.get("phone", "")),
        linkedin=str(data.get("linkedin", "")),
    )
    if data.get("summary"):
        profile.summary = {language: str(data["summary"])}
    skills = [item.strip() for item in re.split(r"[,;\n]", str(data.get("skills", ""))) if item.strip()]
    if skills:
        profile.skills.append(
            SkillGroup(key="skills", label={language: "Skills"}, items=skills)
        )
    return vocabulary.refresh(profile)
