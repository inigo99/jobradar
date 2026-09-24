"""Excel export — one formatted sheet, plus a second sheet of run history."""

from __future__ import annotations

from pathlib import Path

from ..errors import MissingDependencyError
from ..storage import Database
from .csv_export import COLUMNS, export_failure, rows


def export_excel(database: Database, path: Path) -> Path:
    """Write an .xlsx with frozen headers, an auto-filter and sane widths."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise MissingDependencyError(
            "Excel export needs the 'excel' extra.",
            hint="pip install 'jobradar[excel]', or export to CSV with --format csv.",
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        sheet = workbook.create_sheet()
    sheet.title = "Jobs"
    sheet.append([column.replace("_", " ").title() for column in COLUMNS])

    header_fill = PatternFill("solid", fgColor="14625A")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows(database):
        sheet.append([row[column] for column in COLUMNS])

    widths = {"company": 26, "title": 42, "location": 24, "url": 46, "strengths": 40,
              "gaps": 34, "alerts": 44, "notes": 34, "id": 22}
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 15)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    history = workbook.create_sheet("Runs")
    history.append(["Started", "Finished", "Sources", "Fetched", "After dedupe", "Kept", "New", "Errors"])
    for cell in history[1]:
        cell.font = Font(bold=True)
    for run in database.recent_runs(100):
        history.append([
            run.started_at.isoformat(timespec="minutes") if run.started_at else "",
            run.finished_at.isoformat(timespec="minutes") if run.finished_at else "",
            ", ".join(run.sources or []), run.fetched, run.after_dedupe, run.kept, run.new,
            "; ".join(run.errors or []),
        ])
    for index, width in enumerate([20, 20, 40, 10, 14, 8, 8, 50], start=1):
        history.column_dimensions[get_column_letter(index)].width = width

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    except OSError as exc:
        raise export_failure(path, exc) from exc
    return path
